"""Quizni Word (.docx) hujjati sifatida eksport qilish -
[[quiz_pdf_export]]ning DOCX varianti (2026-09-29, foydalanuvchi so'rovi:
o'qituvchilar matnni tahrirlash uchun Word formatini ham xohlashi mumkin).
Xuddi shu tuzilma: bo'sh savollar, so'ng alohida "Javoblar kaliti" bo'limi
(sahifa uzilishi bilan ajratilgan)."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.models.quiz import Question


def _option_letter(index: int) -> str:
    return chr(ord("A") + index)


def build_quiz_docx(*, quiz_name: str, questions: list[Question]) -> bytes:
    """`questions` DB tartibida (ID bo'yicha) keladi deb kutiladi - xuddi
    [[quiz_pdf_export.build_quiz_pdf]]dagi kabi, tasodifiy aralashtirish
    YO'Q (bir nechta nusxa chop etilsa ham bir xil tartib)."""
    document = Document()

    title = document.add_heading(quiz_name, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, question in enumerate(questions, start=1):
        q_paragraph = document.add_paragraph()
        q_run = q_paragraph.add_run(f"{i}. {question.question_text}")
        q_run.bold = True
        q_run.font.size = Pt(12)

        for option_index, option_text in enumerate(question.options):
            option_paragraph = document.add_paragraph(
                f"{_option_letter(option_index)}) {option_text}",
                style="List Bullet",
            )
            option_paragraph.paragraph_format.left_indent = Pt(18)

    document.add_page_break()
    answer_title = document.add_heading("Javoblar kaliti", level=1)
    answer_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for i, question in enumerate(questions, start=1):
        letter = _option_letter(question.correct_option_index)
        document.add_paragraph(f"{i} — {letter}")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
